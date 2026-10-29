import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=120, bottom=120, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_header_styled(doc, text, level=1):
    h = doc.add_heading(level=level)
    run = h.add_run(text)
    run.font.name = 'Arial'
    if level == 1:
        run.font.size = Pt(14)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Navy
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
    elif level == 2:
        run.font.size = Pt(12)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x03, 0x69, 0xA1) # Ocean Blue
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
    elif level == 3:
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A) # Slate Dark
        h.paragraph_format.space_before = Pt(8)
        h.paragraph_format.space_after = Pt(3)
    return h

def add_p(doc, text, bold=False, italic=False, color=(0x33, 0x41, 0x55), space_after=5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = 'Arial'
    run.font.size = Pt(10.5)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = RGBColor(*color)
    p.paragraph_format.line_spacing = 1.15
    p.paragraph_format.space_after = Pt(space_after)
    return p

def create_document():
    doc = docx.Document()

    # Set Margins (Normal 1 inch)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Title / Kop Dokumen
    p_kop = doc.add_paragraph()
    p_kop.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_kop1 = p_kop.add_run("KOMISI PEMILIHAN UMUM KABUPATEN NGAWI\n")
    r_kop1.font.name = 'Arial'
    r_kop1.font.size = Pt(13)
    r_kop1.bold = True
    r_kop1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    r_kop2 = p_kop.add_run("CATATAN TEKNIS & ANALISIS YURIDIS PENGATURAN KELOMPOK KERJA (POKJA)\n")
    r_kop2.font.name = 'Arial'
    r_kop2.font.size = Pt(14)
    r_kop2.bold = True
    r_kop2.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)

    r_kop3 = p_kop.add_run("Pedoman Standar Pembentukan, Pembatasan Kuota Satker, Alokasi Honorarium, dan Akuntabilitas SPJ\nPenyelenggaraan Pemilihan Bupati dan Wakil Bupati Ngawi\n")
    r_kop3.font.name = 'Arial'
    r_kop3.font.size = Pt(10)
    r_kop3.italic = True
    r_kop3.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
    p_kop.paragraph_format.space_after = Pt(14)

    # Horizontal Divider Line (Table 1x1 as a bar)
    t_bar = doc.add_table(rows=1, cols=1)
    t_bar.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_bar = t_bar.rows[0].cells[0]
    c_bar.width = Inches(6.5)
    set_cell_background(c_bar, "1E3A8A")
    p_bar = c_bar.paragraphs[0]
    p_bar.paragraph_format.space_before = Pt(2)
    p_bar.paragraph_format.space_after = Pt(2)
    r_bar = p_bar.add_run(" SISTEM INFORMASI PERENCANAAN & ANGGARAN BADAN AD HOC (SI-PRABU) ")
    r_bar.font.name = 'Arial'
    r_bar.font.size = Pt(9)
    r_bar.bold = True
    r_bar.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # BAB I: DASAR HUKUM
    add_header_styled(doc, "I. DASAR HUKUM & RUJUKAN YURIDIS UTAMA", level=1)
    dasar_hukum = [
        ("1.", "Undang-Undang Nomor 1 Tahun 2015 jo. Undang-Undang Nomor 6 Tahun 2020 tentang Pemilihan Gubernur, Bupati, dan Walikota."),
        ("2.", "Peraturan Menteri Keuangan Republik Indonesia tentang Standar Biaya Masukan (SBM) yang berlaku untuk Tahun Anggaran berjalan."),
        ("3.", "Surat Menteri Keuangan Republik Indonesia Nomor S-647/MK.02/2022 tanggal 5 Agustus 2022 hal Satuan Biaya Masukan Lainnya (SBML) Tahapan Pemilihan Umum dan Tahapan Pemilihan."),
        ("4.", "Keputusan Komisi Pemilihan Umum Nomor 472 Tahun 2022 tentang Satuan Biaya Masukan Lainnya (SBML) di Lingkungan KPU, KPU Provinsi, dan KPU Kabupaten/Kota dalam rangka Tahapan Pemilu dan Pemilihan Tahun 2024."),
        ("5.", "Keputusan Komisi Pemilihan Umum Nomor 543 Tahun 2022 tentang Standar dan Petunjuk Teknis Penyusunan Anggaran Kebutuhan Barang/Jasa dan Honorarium Penyelenggaraan Pemilihan Gubernur dan Wakil Gubernur, Bupati dan Wakil Bupati, dan/atau Wali Kota dan Wakil Wali Kota."),
        ("6.", "Keputusan Komisi Pemilihan Umum Nomor 67 Tahun 2023 tentang Pedoman Teknis Pembentukan dan Pertanggungjawaban Tim Pelaksana Kegiatan/Kelompok Kerja di Lingkungan KPU, KPU Provinsi, dan KPU Kabupaten/Kota."),
        ("7.", "Keputusan Gubernur Jawa Timur Nomor 188/87/KPTS/013/2022 tentang Komponen Pendanaan Bersama Antara Pemerintah Provinsi Jawa Timur dan Pemerintah Kabupaten/Kota dalam Penyelenggaraan Pilkada Serentak Tahun 2024.")
    ]
    for num, txt in dasar_hukum:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(num + " ")
        r1.bold = True
        r1.font.name = 'Arial'
        r1.font.size = Pt(10)
        r2 = p.add_run(txt)
        r2.font.name = 'Arial'
        r2.font.size = Pt(10)

    # BAB II: HAKIKAT & KRITERIA PEMBENTUKAN POKJA
    add_header_styled(doc, "II. HAKIKAT DAN KRITERIA BAKU PEMBENTUKAN POKJA", level=1)
    add_p(doc, "Berdasarkan Lampiran V Bab II Bagian C angka 1 Keputusan KPU No. 543 Tahun 2022 (halaman 65-66), pembentukan Kelompok Kerja (Pokja) harus memenuhi prinsip efisiensi, akuntabilitas, dan 5 kriteria kumulatif:")
    
    kriteria = [
        ("Mempunyai Keluaran (Output) yang Jelas dan Terukur", "Setiap Pokja dibentuk bukan untuk kegiatan seremonial, melainkan wajib menghasilkan output fisik atau yuridis yang terukur (Keputusan, Juknis, SOP, Rekapitulasi Resmi, atau Laporan Akhir)."),
        ("Bersifat Koordinatif Lintas Instansi", "Mengharuskan keterlibatan instansi vertikal atau perangkat daerah eksternal di luar KPU (Polres, Kodim, Bakesbangpol, Dispendukcapil, Satpol PP, Dinkes, Kejaksaan, dsb)."),
        ("Bersifat Temporer dan Insidentil", "Tidak boleh bersifat permanen atau berkelanjutan sepanjang tahun. Pelaksanaannya diprioritaskan di luar jam kerja atau pada masa puncak tahapan tertentu."),
        ("Perangkapan Fungsi / Tugas Tambahan", "Merupakan penugasan tambahan bagi Pejabat Negara, ASN (PNS/PPPK), dan tenaga profesional di samping tugas pokok sehari-hari."),
        ("Dilakukan Secara Selektif, Efektif, dan Efisien", "Dilarang membentuk Pokja untuk tugas-tugas operasional rutin kesekretariatan yang telah melekat pada tugas pokok dan fungsi (tupoksi) organik KPU.")
    ]
    for tit, desc in kriteria:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(4)
        r_t = p.add_run("• " + tit + ": ")
        r_t.bold = True
        r_t.font.name = 'Arial'
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
        r_d = p.add_run(desc)
        r_d.font.name = 'Arial'
        r_d.font.size = Pt(10)

    # BAB III: ANALISIS DUA KUOTA PEMBATASAN POKJA
    add_header_styled(doc, "III. ANALISIS YURIDIS DUA BATASAN POKJA (KUOTA SATKER VS KUOTA PERSONEL)", level=1)
    add_p(doc, "Terdapat dua dimensi pengendalian anggaran yang berbeda dan sering menjadi perdebatan audit antara Tim Perencana dengan BPK/Inspektorat, yaitu batasan kuota kelembagaan (Satker) dan batasan kuota penerimaan individu (Personel):")

    add_header_styled(doc, "1. Batasan Kuota Kelembagaan (Satker): MAKSIMAL 5 POKJA BERHONOR", level=2)
    add_p(doc, "Ketentuan ini diatur secara eksplisit dalam Keputusan KPU Nomor 472 Tahun 2022 Diktum KEEMPAT huruf a angka 5 yang berbunyi:")
    
    # Quote box
    t_q1 = doc.add_table(rows=1, cols=1)
    t_q1.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_q1 = t_q1.rows[0].cells[0]
    c_q1.width = Inches(6.3)
    set_cell_background(c_q1, "F1F5F9")
    set_cell_margins(c_q1, 100, 100, 150, 150)
    p_q1 = c_q1.paragraphs[0]
    p_q1.paragraph_format.space_before = Pt(3)
    p_q1.paragraph_format.space_after = Pt(3)
    r_q1 = p_q1.add_run('"jumlah pokja yang dapat diberikan honorarium yaitu maksimal 5 (lima) pokja masing-masing di Tingkat Pusat, per Provinsi, per Kabupaten/Kota dari seluruh jumlah tahapan pemilu/pemilihan sesuai ketentuan perundang-undangan yang berlaku;"')
    r_q1.font.name = 'Arial'
    r_q1.font.size = Pt(9.5)
    r_q1.italic = True
    r_q1.bold = True
    r_q1.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_p(doc, "Implikasi bagi KPU Kabupaten Ngawi:\n"
               "• Subjek hukum norma ini adalah wilayah teritorial / Satker ('masing-masing per Kabupaten/Kota'), bukan individu.\n"
               "• Meskipun KPT 543/2022 menyediakan katalog 11 jenis Pokja, izin prinsip Kementerian Keuangan (Surat S-647/MK.02/2022) mengunci bahwa Satker KPU Kabupaten/Kota hanya diperkenankan mengalokasikan pembayaran honorarium untuk maksimal 5 Pokja Prioritas selama seluruh rangkaian tahapan Pemilihan.\n"
               "• Jika Satker membentuk Pokja ke-6, ke-7, dan seterusnya, maka Pokja tersebut harus berstatus Tim Kerja Non-Honor (hanya didukung biaya rapat/ATK biasa).")

    add_header_styled(doc, "2. Batasan Kuota Individu (Personel): MAKSIMAL 2 POKJA DALAM 1 TAHUN", level=2)
    add_p(doc, "Ketentuan ini diatur dalam Keputusan KPU Nomor 543 Tahun 2022 Lampiran V Bab II Huruf C angka 1.a (Halaman 66) jo. KPT KPU 472/2022 Diktum Keempat huruf b jo. PMK Standar Biaya Masukan (SBM):")
    
    # Quote box 2
    t_q2 = doc.add_table(rows=1, cols=1)
    t_q2.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_q2 = t_q2.rows[0].cells[0]
    c_q2.width = Inches(6.3)
    set_cell_background(c_q2, "F1F5F9")
    set_cell_margins(c_q2, 100, 100, 150, 150)
    p_q2 = c_q2.paragraphs[0]
    p_q2.paragraph_format.space_before = Pt(3)
    p_q2.paragraph_format.space_after = Pt(3)
    r_q2 = p_q2.add_run('"...jumlah tim yang dapat diberikan honorarium bagi: 1) Pejabat Negara; 2) Pejabat Eselon I; 3) Pejabat Eselon II; 4) Pejabat Eselon III; 5) Pejabat Eselon IV; 6) Pelaksana; dan 7) Pejabat Fungsional, pada tim dimaksud, untuk penerimaan honorarium yang berasal dari pembentukan tim/Pokja paling banyak 2 (dua) tim setiap tahunnya."')
    r_q2.font.name = 'Arial'
    r_q2.font.size = Pt(9.5)
    r_q2.italic = True
    r_q2.bold = True
    r_q2.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    add_p(doc, "Implikasi bagi Personel KPU Kabupaten Ngawi:\n"
               "• Subjek hukum norma ini adalah person (Komisioner, Pejabat Struktural, Staf ASN/PPPK).\n"
               "• Tidak boleh ada satupun personel yang namanya tercantum dan menerima honorarium lebih dari 2 Pokja dalam satu tahun anggaran yang sama.\n"
               "• Jika seseorang menerima honor Pokja ke-3 dalam tahun berjalan, pembayaran tersebut berstatus Kelebihan Bayar / Kerugian Keuangan Daerah yang wajib disetorkan kembali ke Kas Daerah (TGR).")

    # BAB IV: MATRIKS RINGKASAN
    add_header_styled(doc, "IV. MATRIKS PERBANDINGAN NORMA & PENGENDALIAN POKJA", level=1)
    
    t_mat = doc.add_table(rows=8, cols=3)
    t_mat.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_mat = ["Norma / Parameter", "Batasan Maksimal (Aturan Ketat)", "Dasar Regulasi Resmi"]
    widths_mat = [Inches(2.0), Inches(2.6), Inches(1.9)]
    
    for c_idx, h_text in enumerate(headers_mat):
        c = t_mat.rows[0].cells[c_idx]
        c.width = widths_mat[c_idx]
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 100, 100, 120, 120)
        p = c.paragraphs[0]
        r = p.add_run(h_text)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_mat = [
        ("Kuota Pokja Satker KPU Kab/Kota", "Maksimal 5 (lima) Pokja yang diberikan honorarium dari seluruh tahapan Pemilihan.", "KPT KPU 472/2022 Diktum IV.a.5"),
        ("Kuota Penerimaan per Personel", "Maksimal 2 (dua) Pokja dalam 1 Tahun Anggaran per individu (Komisioner/ASN).", "KPT KPU 543/2022 Hal. 66 jo. PMK SBM"),
        ("Durasi Masa Kerja Pokja", "Paling lama 3 (tiga) Bulan berturut-turut.", "KPT KPU 543/2022 Lamp. IV & Hal. 66"),
        ("Komposisi Personel Internal KPU", "Paling banyak 10 orang dari internal KPU Kab/Kota (sisanya wajib instansi luar/Pemda).", "KPT KPU 543/2022 Hal. 70 jo. KPT 472/2022"),
        ("Larangan Irisan (Double Funding)", "Dilarang keras menerima duplikasi honor jika tahapan Pemilu (APBN) dan Pilkada beririsan.", "KPT KPU 472/2022 Diktum VI"),
        ("Larangan Pegawai Outsource", "Tenaga pendukung outsource/kontrak dilarang keras menerima honor Pokja.", "KPT KPU 543/2022 Hal. 37 & Hal. 70"),
        ("Keluaran SPJ Wajib", "Laporan Hasil Kerja Lengkap (Bab I s.d. Bab IV) yang disahkan oleh KPA.", "KPT KPU 543/2022 Hal. 72–75")
    ]

    for r_idx, row_data in enumerate(data_mat):
        row = t_mat.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate(row_data):
            c = row.cells[c_idx]
            c.width = widths_mat[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 80, 80, 100, 100)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            if c_idx == 0:
                r.bold = True
                r.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            elif c_idx == 1 and "Maksimal" in val:
                r.bold = True
                r.font.color.rgb = RGBColor(0xB9, 0x1C, 0x1C)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # BAB V: KATALOG 11 POKJA & REKOMENDASI 5 PRIORITAS
    add_header_styled(doc, "V. KATALOG 11 POKJA RESMI & REKOMENDASI 5 POKJA PRIORITAS NGAWI", level=1)
    add_p(doc, "Berdasarkan Lampiran IV dan Lampiran V halaman 68-69 KPT 543/2022, terdapat 11 jenis Pokja resmi di KPU Kabupaten/Kota:")

    t_pok = doc.add_table(rows=12, cols=5)
    t_pok.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_pok = ["No", "Nama Kelompok Kerja (Pokja)", "Maks Internal", "Total Kuota", "Status Rekomendasi Ngawi"]
    w_pok = [Inches(0.4), Inches(2.9), Inches(1.0), Inches(0.9), Inches(1.3)]
    for c_idx, txt in enumerate(h_pok):
        c = t_pok.rows[0].cells[c_idx]
        c.width = w_pok[c_idx]
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 90, 90, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(txt)
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_pok = [
        ("1", "Pokja Penyusunan, Perencanaan, Perubahan/Revisi Rencana Anggaran Pemilihan", "10 orang", "13 orang", "PRIORITAS 1 (Berhonor)"),
        ("2", "Pokja Penyusunan Produk-Produk Hukum (Keputusan)", "10 orang", "13 orang", "Tim Kerja Kedinasan"),
        ("3", "Pokja Fasilitasi Pembentukan Badan Ad Hoc", "10 orang", "15 orang", "Tim Kerja Kedinasan"),
        ("4", "Pokja Pemutakhiran Data Pemilih dan Penyusunan Daftar Pemilih", "10 orang", "15 orang", "PRIORITAS 2 (Berhonor)"),
        ("5", "Pokja Verifikasi Persyaratan Pasangan Calon Perseorangan", "10 orang", "15 orang", "Kondisional (Jika Ada Calon)"),
        ("6", "Pokja Pendaftaran dan Penetapan Pasangan Calon", "10 orang", "13 orang", "PRIORITAS 3 (Berhonor)"),
        ("7", "Pokja Sosialisasi dan Partisipasi Masyarakat / Bimtek", "10 orang", "15 orang", "Tim Kerja Kedinasan"),
        ("8", "Pokja Pelaksanaan Kampanye serta Pelayanan Dana Kampanye", "10 orang", "15 orang", "Tim Kerja Kedinasan"),
        ("9", "Pokja Perencanaan, Pengadaan dan Pendistribusian Logistik", "10 orang", "13 orang", "PRIORITAS 4 (Berhonor)"),
        ("10", "Pokja Pemungutan & Penghitungan Suara serta Rekapitulasi", "10 orang", "15 orang", "PRIORITAS 5 (Berhonor)"),
        ("11", "Pokja Penyelesaian Sengketa Hukum", "10 orang", "13 orang", "Kondisional (Jika Ada Gugatan)")
    ]

    for r_idx, row_d in enumerate(data_pok):
        row = t_pok.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate(row_d):
            c = row.cells[c_idx]
            c.width = w_pok[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(8.5)
            if c_idx == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r.bold = True
            elif c_idx == 4 and "PRIORITAS" in val:
                r.bold = True
                r.font.color.rgb = RGBColor(0x04, 0x78, 0x57) # Green
            elif c_idx == 4 and "Kondisional" in val:
                r.italic = True
                r.font.color.rgb = RGBColor(0xB4, 0x53, 0x09)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # BAB VI: STANDAR HONORARIUM POKJA
    add_header_styled(doc, "VI. STANDAR SATUAN BIAYA HONORARIUM POKJA KPU KABUPATEN/KOTA", level=1)
    add_p(doc, "Berdasarkan Lampiran II Keputusan KPU Nomor 472 Tahun 2022 (SBML Tahapan Pemilihan Bupati/Walikota), besaran honorarium Pokja KPU Kab/Kota per bulan (Orang/Bulan - OB) adalah:")

    t_hon = doc.add_table(rows=6, cols=4)
    t_hon.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_hon = ["Kedudukan dalam Pokja", "Besaran SBML (Rp/OB)", "Ketentuan Personel", "Maksimal Durasi"]
    w_hon = [Inches(2.0), Inches(1.5), Inches(2.0), Inches(1.0)]
    for c_idx, txt in enumerate(h_hon):
        c = t_hon.rows[0].cells[c_idx]
        c.width = w_hon[c_idx]
        set_cell_background(c, "1E3A8A")
        set_cell_margins(c, 90, 90, 100, 100)
        p = c.paragraphs[0]
        r = p.add_run(txt)
        r.font.name = 'Arial'
        r.font.size = Pt(9)
        r.bold = True
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    data_hon = [
        ("Pengarah", "Rp 2.000.000", "Ketua/Anggota KPU yang membidangi (Maks 2 orang)", "3 Bulan"),
        ("Penanggung Jawab", "Rp 1.800.000", "Ketua/Anggota KPU yang membidangi (1 orang)", "3 Bulan"),
        ("Ketua", "Rp 1.500.000", "Anggota KPU atau Sekretaris KPU (1 orang)", "3 Bulan"),
        ("Sekretaris", "Rp 1.300.000", "Sekretaris KPU atau Kasubbag yang membidangi (1 orang)", "3 Bulan"),
        ("Anggota", "Rp 1.200.000", "Pejabat Fungsional, ASN KPU & Perwakilan Eksternal Pemda", "3 Bulan")
    ]

    for r_idx, row_d in enumerate(data_hon):
        row = t_hon.rows[r_idx + 1]
        bg = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
        for c_idx, val in enumerate(row_d):
            c = row.cells[c_idx]
            c.width = w_hon[c_idx]
            set_cell_background(c, bg)
            set_cell_margins(c, 70, 70, 90, 90)
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(val)
            r.font.name = 'Arial'
            r.font.size = Pt(9)
            if c_idx == 0:
                r.bold = True
            elif c_idx == 1:
                r.bold = True
                r.font.color.rgb = RGBColor(0x04, 0x78, 0x57)
            elif c_idx == 3:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # BAB VII: STANDAR SPJ & SISTEMATIKA LAPORAN BAB I-IV
    add_header_styled(doc, "VII. STANDAR PERTANGGUNGJAWABAN (SPJ) DAN SISTEMATIKA LAPORAN", level=1)
    add_p(doc, "Sesuai Lampiran V halaman 72-75 KPT 543/2022, pencairan honorarium Pokja (Akun 521213) bukan merupakan hak otomatis bulanan, melainkan mensyaratkan 6 bukti dukung wajib:")
    
    spj_items = [
        ("Surat Keputusan (SK) Pembentukan Pokja", "Diterbitkan oleh KPA / Ketua KPU dengan mencantumkan nama, jabatan kedinasan, kedudukan dalam Pokja, dan masa tugas."),
        ("Surat Tugas", "Diterbitkan oleh KPA untuk penugasan rapat dan koordinasi di lapangan."),
        ("Undangan Rapat Koordinasi", "Disampaikan kepada seluruh anggota Pokja dari unsur internal maupun eksternal."),
        ("Daftar Hadir Asli", "Presensi kehadiran fisik bertandatangan basah pada setiap rapat pelaksanaan tugas."),
        ("Risalah / Notulensi Rapat", "Mencatat secara sistematis tanggal, waktu, pokok bahasan, dan kesimpulan keputusan rapat."),
        ("Laporan Hasil Pelaksanaan Tugas Pokja (Wajib Bab I s.d. Bab IV)", "Merupakan bukti output riil yang disahkan oleh KPA dengan sistematika baku:")
    ]
    for num, (t_s, d_s) in enumerate(spj_items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run(f"{num}. {t_s}: ")
        r1.bold = True
        r1.font.name = 'Arial'
        r1.font.size = Pt(9.5)
        r2 = p.add_run(d_s)
        r2.font.name = 'Arial'
        r2.font.size = Pt(9.5)

    add_p(doc, "Sistematika Baku Laporan Output Pokja (KPT 543/2022 Halaman 73-75):\n"
               "• BAB I PENDAHULUAN: Latar Belakang (isu strategis), Sasaran, Maksud & Tujuan.\n"
               "• BAB II GAMBARAN UMUM: Dasar Hukum, Jangka Waktu Pelaksanaan, Rencana Anggaran Biaya (RAB), Susunan Personel & Uraian Tugas.\n"
               "• BAB III HASIL TIM PELAKSANAAN KEGIATAN: Langkah Program/Kegiatan, Output/Keluaran Terukur disertai Data Valid.\n"
               "• BAB IV PENUTUP: Kesimpulan Capaian Pokja & Rekomendasi/Tindak Lanjut ke Depan.\n"
               "• LAMPIRAN: Fotokopi SK Pokja, Surat Tugas, Undangan, Notulensi, Presensi Hadir, dan Dokumentasi Foto Kegiatan.")

    # BAB VIII: REKOMENDASI AUDIT UNTUK KPU KABUPATEN NGAWI
    add_header_styled(doc, "VIII. REKOMENDASI STRATEGIS & MITIGASI RISIKO AUDIT (KPU NGAWI)", level=1)
    rek_items = [
        "Pengendalian Plafon 5 Pokja: KPA dan Kasubbag Perencanaan KPU Kab. Ngawi menyepakati 5 Pokja Prioritas yang akan dialokasikan honorariumnya dalam Dokumen Pelaksanaan Anggaran (DPA) Hibah NPHD.",
        "Penguncian Personel (Maks 2 Pokja): Membuat matriks silang (cross-check table) seluruh nama pegawai/komisioner untuk memastikan tidak ada personel yang menerima honor lebih dari 2 Pokja dalam tahun anggaran yang sama.",
        "Disiplin Masa Kerja (Maks 3 Bulan): SK Pokja tidak boleh dibuat dengan masa kerja melebihi 3 bulan berturut-turut. Tanggal SK harus menyesuaikan dengan timeline tahapan riil di lapangan.",
        "Pelibatan Instansi Eksternal: Kuota internal KPU wajib dibatasi maksimal 10 orang per Pokja, sedangkan 3-5 personel lainnya dialokasikan untuk mitra kerja terkait (Polres, Kodim, Dispendukcapil, Satpol PP, dsb).",
        "Validasi SPJ Sebelum Pembayaran: Bendahara Pengeluaran Pembantu (BPP) dilarang menerbitkan bukti bayar honorarium sebelum Laporan Hasil Pelaksanaan Tugas Pokja (Bab I-IV) disahkan secara resmi oleh KPA."
    ]
    for r_k in rek_items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("✔ " + r_k)
        r.font.name = 'Arial'
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

    # Footer Sign
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    p_sign = doc.add_paragraph()
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_s1 = p_sign.add_run("Ngawi, Oktober 2026\n")
    r_s1.font.name = 'Arial'
    r_s1.font.size = Pt(10)
    r_s2 = p_sign.add_run("Disusun Oleh:\nTim Sistem Perencanaan Anggaran Badan Ad Hoc (SI-PRABU)\nKomisi Pemilihan Umum Kabupaten Ngawi\n\n\n")
    r_s2.font.name = 'Arial'
    r_s2.font.size = Pt(10)
    r_s2.italic = True
    r_s3 = p_sign.add_run("LEAD DEVELOPER & SYSTEM ARCHITECT (emenje)")
    r_s3.font.name = 'Arial'
    r_s3.font.size = Pt(10.5)
    r_s3.bold = True
    r_s3.underline = True

    # Save to public/docs
    out_path = os.path.join("public", "docs", "Catatan_Analisis_Regulasi_Pokja_KPU_Ngawi.docx")
    doc.save(out_path)
    print(f"Document successfully created at: {out_path}")

if __name__ == '__main__':
    create_document()
